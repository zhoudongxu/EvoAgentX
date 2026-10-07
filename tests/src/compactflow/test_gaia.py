import asyncio, io, json, zipfile
import pytest
from evoagentx.compactflow.gaia import (
    GaiaToolSession,
    GaiaUnavailable,
    GaiaBudgetExceeded,
    attach_assets,
    initialize_gaia_run,
)
from evoagentx.compactflow.gaia_tools import GaiaBackends
from evoagentx.compactflow.benchmarks import (
    BenchmarkTask,
    import_dataset,
    assign_exact_partitions,
)
from evoagentx.compactflow.baseline_native import ModelSession
from evoagentx.compactflow.baseline_controls import PolicyBaselineAdapter
from evoagentx.compactflow.baselines import (
    AFlowAdapter,
    EvoAgentXAdapter,
    ConstructionBaselineRunner,
)
from evoagentx.compactflow.paper_workflow import validate_spec
from evoagentx.compactflow.llm import HashEmbedder
from test_baselines import config, manifest, FakeModelClient


def gaia_config(tmp_path):
    c = config()
    c["evaluation"]["generation_seeds"] = [42]
    c["tools"]["gaia"].update(asset_root=str(tmp_path / "raw"), replay_time_scale=0)
    (tmp_path / "raw").mkdir(exist_ok=True)
    return c


def gaia_tasks(tmp_path):
    tasks = []
    for split in ("source", "validation", "target"):
        name = split + ".txt"
        (tmp_path / "raw" / name).write_text("One item for " + split)
        tasks.append(
            BenchmarkTask(
                "GAIA",
                split,
                "Count the " + split + " items.",
                "SECRET_GOLD",
                split,
                attachments=[name],
                split=split,
                metadata={"level": 2, "validation_fold": 1}
                if split == "validation"
                else {"level": 1},
            )
        )
    attach_assets(tasks, tmp_path / "raw")
    return tasks


class FakeBackend:
    calls = []

    def __init__(self, *args):
        pass

    async def estimate(self, tool, args):
        return 0

    async def call(self, tool, args):
        self.calls.append((tool, args))
        return {"text": "There is one item."}


class FakeGaiaClient(FakeModelClient):
    async def text(self, system, prompt, *, component, seed):
        previous = await super().text(system, prompt, component=component, seed=seed)
        if component == "executor":
            data = json.loads(prompt)
            if not data["observations"]:
                return json.dumps(
                    {
                        "tool_calls": [
                            {
                                "tool": "read_file",
                                "arguments": {
                                    "asset_id": data["task"]["assets"][0]["asset_id"]
                                },
                            }
                        ]
                    }
                )
            return json.dumps({"final": {k: "1" for k in data["output_fields"]}})
        if "<operator_description>" in prompt:
            return previous.replace("operator.Custom", "operator.GAIAToolAgent")
        if component == "query":
            return json.dumps({"query": "solve with evidence"})
        if component == "selector":
            return json.dumps({"selected": []})
        try:
            data = json.loads(prompt)
        except ValueError:
            return previous
        return json.dumps(
            {
                "nodes": [
                    {
                        "id": "solve",
                        "tool": "gaia_agent",
                        "instruction": "Solve",
                        "inputs": {
                            "question": "$input.question",
                            "assets": "$input.assets",
                        },
                        "outputs": ["answer"],
                    }
                ],
                "sinks": ["solve"],
                "applied_policies": [],
                "unapplied_policies": [
                    {"id": p["id"], "reason": "not applicable"}
                    for p in data.get("selected_policies", [])
                ],
            }
        )


def ready(c, **kwargs):
    return {
        "status": "complete",
        "reasons": [],
        "models": c["tools"]["gaia"]["models"],
        "packages": {},
        "tool_schema_digest": "test",
    }


@pytest.mark.parametrize(
    "method",
    [
        "aflow",
        "evoagentx",
        "base_planner",
        "expert_policies",
        "static_library",
        "compactflow",
    ],
)
def test_six_methods_share_gaia_tools_and_frozen_target(tmp_path, monkeypatch, method):
    import evoagentx.compactflow.gaia as module
    import evoagentx.compactflow.gaia_tools as backends

    monkeypatch.setattr(module, "capability_report", ready)
    monkeypatch.setattr(backends, "GaiaBackends", FakeBackend)
    c = gaia_config(tmp_path)
    tasks = gaia_tasks(tmp_path)
    m = manifest(tmp_path, tasks)
    evo = tmp_path / "evo"
    (evo / "policy_snapshots").mkdir(parents=True)
    for name in ("initial", "bootstrap", "final"):
        (evo / "policy_snapshots" / f"{name}.json").write_text(
            '{"policies":[],"evidence":[]}'
        )
    (evo / "checkpoint.json").write_text('{"stage":"complete"}')
    before = {p.name: p.read_bytes() for p in (evo / "policy_snapshots").glob("*")}

    def adapter():
        if method == "aflow":
            return AFlowAdapter(client_factory=FakeGaiaClient)
        if method == "evoagentx":
            return EvoAgentXAdapter(client_factory=FakeGaiaClient)
        return PolicyBaselineAdapter(
            method,
            evolution_dir=evo,
            client_factory=FakeGaiaClient,
            embedder=HashEmbedder(),
        )

    def runner():
        return ConstructionBaselineRunner(
            tasks,
            output=tmp_path / "out",
            adapters=[adapter()],
            config=c,
            manifest=m,
            evolution_dir=evo,
            evaluate_fn=lambda t, a: {"quality": float(a == "1")},
        )

    FakeBackend.calls = []
    FakeModelClient.requests = []
    summary = asyncio.run(runner().run())
    assert summary["methods"][method]["GAIA"]["status"] == "complete", summary[
        "methods"
    ][method]["GAIA"]
    rows = [
        json.loads(l)
        for l in (tmp_path / "out/baseline_records.jsonl").read_text().splitlines()
    ]
    assert rows and all(
        r["status"] == "complete" and r["quality"] == 1 and r["token_cost"] > 0
        for r in rows
    )
    assert all(r["metadata"]["tool_accounting"]["tool_calls"] > 0 for r in rows)
    assert {
        (r["benchmark"], r["task_id"], r["seed"])
        for r in rows
        if r["split"] == "target"
    } == {("GAIA", "target", 42)}
    assert {
        p.name: p.read_bytes() for p in (evo / "policy_snapshots").glob("*")
    } == before
    assert FakeBackend.calls
    count = len(FakeModelClient.requests)
    tools = len(FakeBackend.calls)
    asyncio.run(runner().run(resume=True))
    assert len(FakeModelClient.requests) == count and len(FakeBackend.calls) == tools


def test_target_tools_require_freeze(tmp_path):
    c = gaia_config(tmp_path)
    t = gaia_tasks(tmp_path)[2]
    with pytest.raises(GaiaUnavailable, match="frozen"):
        GaiaToolSession(
            c,
            t.public_input(),
            output=tmp_path / "out",
            workspace=tmp_path / "call",
            seed=42,
            backend=FakeBackend(),
        )


def test_capture_exact_replay_and_no_live_fallback(tmp_path):
    c = gaia_config(tmp_path)
    t = gaia_tasks(tmp_path)[0]
    FakeBackend.calls = []

    def session():
        return GaiaToolSession(
            c,
            t.public_input(),
            output=tmp_path / "out",
            workspace=tmp_path / "call",
            seed=42,
            backend=FakeBackend(),
        )

    args = {"asset_id": t.public_input()["assets"][0]["asset_id"]}
    asyncio.run(session().invoke("read_file", args))
    assert len(FakeBackend.calls) == 1
    c["tools"]["gaia"]["mode"] = "replay"
    asyncio.run(session().invoke("read_file", args))
    assert len(FakeBackend.calls) == 1
    with pytest.raises(GaiaUnavailable, match="fallback"):
        asyncio.run(session().invoke("search", {"query": "missing request"}))
    c["tools"]["gaia"]["mode"] = "capture"
    t.task_id = "different-task"
    asyncio.run(session().invoke("read_file", args))
    assert len(FakeBackend.calls) == 2


def test_asset_hash_and_private_input_protection(tmp_path):
    c = gaia_config(tmp_path)
    t = gaia_tasks(tmp_path)[0]
    public = t.public_input()
    assert "SECRET_GOLD" not in json.dumps(public)
    with pytest.raises(ValueError, match="non-public"):
        GaiaToolSession(
            c,
            {**public, "answer": "secret"},
            output=tmp_path,
            workspace=tmp_path,
            seed=1,
            backend=FakeBackend(),
        )
    (tmp_path / "raw/source.txt").write_text("tampered")
    with pytest.raises(GaiaUnavailable, match="checksum"):
        GaiaBackends(c, public, tmp_path / "assets")


def test_model_cancellation_never_reports_zero_success(tmp_path):
    class Cancel(FakeModelClient):
        async def text(self, *a, **k):
            raise asyncio.CancelledError()

    c = gaia_config(tmp_path)

    async def run():
        session = ModelSession(c, 42, tmp_path / "model", client_factory=Cancel)
        with pytest.raises(asyncio.CancelledError):
            await session.text([{"role": "user", "content": "test"}], "executor", "x")
        assert not session.accounting()["usage_complete"]

    asyncio.run(run())


def test_actual_aux_usage_and_replay_budget(tmp_path):
    class Vision(FakeBackend):
        async def estimate(self, *a):
            return 20

        async def call(self, *a):
            return {
                "text": "one",
                "_usage": {
                    "prompt_tokens": 7,
                    "completion_tokens": 3,
                    "total_tokens": 10,
                    "usage_complete": True,
                },
            }

    c = gaia_config(tmp_path)
    t = gaia_tasks(tmp_path)[0]

    async def run():
        owner = ModelSession(c, 42, tmp_path / "models", client_factory=FakeModelClient)
        session = GaiaToolSession(
            c,
            t.public_input(),
            output=tmp_path / "out",
            workspace=tmp_path / "call",
            seed=42,
            backend=Vision(),
            model_session=owner,
        )
        await session.invoke("vision", {"asset_id": "fake", "question": "count"})
        assert (
            owner.accounting()["total_tokens"] == 10
            and owner.used == 10
            and owner.reserved == 0
        )
        await session.invoke("vision", {"asset_id": "fake", "question": "count"})
        assert (
            owner.accounting()["total_tokens"] == 20
            and session.accounting()["physical_vision_tokens"] == 10
        )

    asyncio.run(run())


def test_exact_reference_split_and_level_import():
    rows = []
    for i in range(150):
        word = chr(97 + i // 26) + chr(97 + i % 26)
        rows.append(
            {
                "task_id": str(i),
                "Question": "Explain " + word,
                "Final answer": "x",
                "Level": str(i % 3 + 1),
                "file_name": "",
            }
        )
    tasks = import_dataset("GAIA", rows, revision="pinned", original_split="validation")
    assign_exact_partitions(
        tasks,
        seed=42,
        source_count=90,
        validation_count=30,
        target_count=30,
        validation_folds=5,
    )
    assert [
        sum(t.split == s for t in tasks) for s in ("source", "validation", "target")
    ] == [90, 30, 30]
    assert {t.metadata["level"] for t in tasks} == {1, 2, 3}
    folds = [
        {t.task_id for t in tasks if t.metadata.get("validation_fold") == i}
        for i in range(1, 6)
    ]
    assert all(len(f) == 6 for f in folds) and len(set.union(*folds)) == 30


def test_zip_traversal_and_tool_budget(tmp_path):
    c = gaia_config(tmp_path)
    t = gaia_tasks(tmp_path)[0]
    b = GaiaBackends(c, t.public_input(), tmp_path / "assets")
    z = io.BytesIO()
    with zipfile.ZipFile(z, "w") as f:
        f.writestr("../secret.txt", "bad")
    asset = b.save_asset(z.getvalue(), "bad.zip")
    with pytest.raises(ValueError, match="unsafe ZIP"):
        b.tool_unpack_zip(asset["asset_id"])
    c["tools"]["gaia"]["max_tool_calls"] = 1

    async def run():
        s = GaiaToolSession(
            c,
            t.public_input(),
            output=tmp_path / "out",
            workspace=tmp_path / "call",
            seed=42,
            backend=FakeBackend(),
        )
        await s.invoke("read_file", {})
        with pytest.raises(GaiaBudgetExceeded):
            await s.invoke("read_file", {})

    asyncio.run(run())


@pytest.mark.parametrize(
    "ext", [".txt", ".csv", ".jsonld", ".xml", ".pdb", ".py", ".xlsx", ".docx", ".pptx"]
)
def test_document_and_table_formats(tmp_path, ext):
    c = gaia_config(tmp_path)
    path = tmp_path / "raw" / ("file" + ext)
    if ext == ".xlsx":
        import openpyxl

        book = openpyxl.Workbook()
        book.active.append(["hello", 7])
        book.save(path)
    elif ext == ".docx":
        from docx import Document

        doc = Document()
        doc.add_paragraph("hello")
        doc.save(path)
    elif ext == ".pptx":
        from pptx import Presentation

        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[0])
        slide.shapes.title.text = "hello"
        deck.save(path)
    else:
        path.write_text("hello")
    task = BenchmarkTask(
        "GAIA", "x", "Read", "secret", "x", attachments=[path.name], split="source"
    )
    attach_assets([task], tmp_path / "raw")
    b = GaiaBackends(c, task.public_input(), tmp_path / "assets")
    assert (
        "hello"
        in b.tool_read_file(task.public_input()["assets"][0]["asset_id"])["text"]
    )


def test_typed_gaia_tool_rejected_for_other_benchmarks(tmp_path):
    spec = {
        "nodes": [
            {
                "id": "s",
                "tool": "gaia_agent",
                "instruction": "solve",
                "inputs": {"q": "$input.question"},
                "outputs": ["answer"],
            }
        ],
        "sinks": ["s"],
        "applied_policies": [],
        "unapplied_policies": [],
    }
    with pytest.raises(ValueError, match="unregistered"):
        validate_spec(spec, {"benchmark": "MATH", "question": "q"}, set())


def test_nested_legacy_xls_and_xml(tmp_path):
    xlwt = pytest.importorskip("xlwt")
    config = gaia_config(tmp_path)
    task = gaia_tasks(tmp_path)[0]
    book = xlwt.Workbook()
    sheet = book.add_sheet("categories")
    sheet.write(0, 0, "hello")
    sheet.write(0, 1, 7)
    data = io.BytesIO()
    book.save(data)
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as z:
        z.writestr("food.xls", data.getvalue())
        z.writestr("CATEGORIES.xml", "<items><item>hello</item></items>")
    backend = GaiaBackends(config, task.public_input(), tmp_path / "assets")
    asset = backend.save_asset(archive.getvalue(), "fixture.zip")
    unpacked = backend.tool_unpack_zip(asset["asset_id"])["assets"]
    assert len(unpacked) == 2
    for entry in unpacked:
        assert "hello" in backend.tool_read_file(entry["asset_id"])["text"]


def test_tool_backend_global_concurrency_bound(tmp_path):
    import threading, time

    config = gaia_config(tmp_path)
    task = gaia_tasks(tmp_path)[0]
    backend = GaiaBackends(config, task.public_input(), tmp_path / "assets")
    state = {"active": 0, "peak": 0}
    lock = threading.Lock()

    def delayed(query):
        with lock:
            state["active"] += 1
            state["peak"] = max(state["peak"], state["active"])
        time.sleep(0.03)
        with lock:
            state["active"] -= 1
        return {"text": query}

    backend.tool_search = delayed

    async def run():
        return await asyncio.gather(
            *(backend.call("search", {"query": str(i)}) for i in range(12))
        )

    assert len(asyncio.run(run())) == 12
    assert state["peak"] <= 4


def test_aux_cancel_durable_unresolved_request(tmp_path):
    class Cancel(FakeBackend):
        async def call(self, *a):
            raise asyncio.CancelledError()

    c = gaia_config(tmp_path)
    task = gaia_tasks(tmp_path)[0]

    async def run():
        s = GaiaToolSession(
            c,
            task.public_input(),
            output=tmp_path / "out",
            workspace=tmp_path / "call",
            seed=42,
            backend=Cancel(),
        )
        with pytest.raises(asyncio.CancelledError):
            await s.invoke("read_file", {})
        assert not s.accounting()["usage_complete"]
        with pytest.raises(GaiaUnavailable, match="unresolved"):
            await s.invoke("read_file", {})

    asyncio.run(run())


def test_sessions_cannot_observe_another_methods_generated_assets(tmp_path):
    c = gaia_config(tmp_path)
    task = gaia_tasks(tmp_path)[0]
    first = GaiaBackends(c, task.public_input(), tmp_path / "assets")
    generated = first.save_asset(b"private observation", "method-one.txt")
    second = GaiaBackends(c, task.public_input(), tmp_path / "assets")
    with pytest.raises(ValueError, match="this session"):
        second.resolve(generated["asset_id"])
    # Only replay of the exact verified observation grants the generated asset.
    second.verify_result_assets({"assets": [generated]})
    assert second.resolve(generated["asset_id"]).read_bytes() == b"private observation"


def test_modified_observation_and_asset_locks_fail_resume(tmp_path, monkeypatch):
    import evoagentx.compactflow.gaia as module

    monkeypatch.setattr(module, "capability_report", ready)
    c = gaia_config(tmp_path)
    tasks = gaia_tasks(tmp_path)
    out = tmp_path / "out"
    initialize_gaia_run(out, c, tasks)
    session = GaiaToolSession(
        c,
        tasks[0].public_input(),
        output=out,
        workspace=out / "call",
        seed=42,
        backend=FakeBackend(),
    )
    asyncio.run(session.invoke("read_file", {}))
    snap = next((out / "tool_snapshots").glob("*.json"))
    original = snap.read_text()
    value = json.loads(original)
    value["observation"]["result"]["text"] = "tampered"
    snap.write_text(json.dumps(value))
    with pytest.raises(ValueError, match="modified"):
        initialize_gaia_run(out, c, tasks)
    snap.write_text(original)
    tasks[0].metadata["gaia_assets"][0]["sha256"] = "changed"
    with pytest.raises(ValueError, match="lock mismatch"):
        initialize_gaia_run(out, c, tasks)


def test_missing_vision_usage_fails_closed(tmp_path):
    c = gaia_config(tmp_path)
    t = gaia_tasks(tmp_path)[0]
    s = GaiaToolSession(
        c,
        t.public_input(),
        output=tmp_path / "out",
        workspace=tmp_path / "call",
        seed=42,
        backend=FakeBackend(),
    )
    with pytest.raises(GaiaUnavailable, match="usage"):
        asyncio.run(s.invoke("vision", {"asset_id": "x", "question": "q"}))
    assert not s.accounting()["usage_complete"]


def test_bounded_malformed_agent_output(tmp_path):
    c = gaia_config(tmp_path)
    c["tools"]["gaia"]["max_agent_steps"] = 2
    t = gaia_tasks(tmp_path)[0]

    class Malformed:
        async def json(self, *args, **kwargs):
            raise ValueError("bad JSON")

    s = GaiaToolSession(
        c,
        t.public_input(),
        output=tmp_path / "out",
        workspace=tmp_path / "call",
        seed=42,
        backend=FakeBackend(),
    )
    with pytest.raises(GaiaBudgetExceeded, match="step budget"):
        asyncio.run(s.run_agent(Malformed(), "solve", {}, ["answer"]))
    assert s.steps == 2


def test_gaia_evolution_real_shared_tools_and_resume(tmp_path, monkeypatch):
    from examples.compactflow.run_evolution import LiveAdapter
    from evoagentx.compactflow.evolution import EvolutionConfig, EvolutionRunner
    from evoagentx.compactflow.policy import PolicyLibrary
    import evoagentx.compactflow.gaia as module
    import evoagentx.compactflow.gaia_tools as backends

    monkeypatch.setattr(module, "capability_report", ready)
    monkeypatch.setattr(backends, "GaiaBackends", FakeBackend)
    c = gaia_config(tmp_path)
    tasks = gaia_tasks(tmp_path)
    for t in tasks:
        t.answer = "1"
    output = tmp_path / "evolution"
    library = PolicyLibrary(embedder=HashEmbedder())
    adapter = LiveAdapter(c, library, output=output, client_factory=FakeGaiaClient)
    ec = EvolutionConfig(
        bootstrap_tasks=0,
        rounds=1,
        source_tasks_per_round=1,
        max_candidates_per_round=1,
        validation_folds=1,
        validation_tasks_per_candidate=1,
        heldout_seeds=(101,),
        target_seeds=(42,),
    )

    def build():
        return EvolutionRunner(
            tasks,
            config=ec,
            library=PolicyLibrary(embedder=HashEmbedder()),
            output=output,
            run_variant=adapter.run_variant,
            distill_candidate=lambda *args: None,
            config_payload=c,
            code_commit="test",
        )

    FakeBackend.calls = []
    FakeModelClient.requests = []
    asyncio.run(build().run())
    rows = [
        json.loads(l)
        for l in (output / "target_records.jsonl").read_text().splitlines()
    ]
    assert len(rows) == 2 and all(row["evidence"]["feedback"]["valid"] for row in rows)
    assert all(row["record"]["tool_accounting"]["tool_calls"] > 0 for row in rows)
    before = (output / "policy_snapshots/final.json").read_bytes()
    count = len(FakeModelClient.requests)
    asyncio.run(build().run(resume=True))
    assert len(FakeModelClient.requests) == count
    assert (output / "policy_snapshots/final.json").read_bytes() == before
