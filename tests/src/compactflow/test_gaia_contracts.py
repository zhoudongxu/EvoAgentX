import asyncio, copy, json
from contextlib import aclosing
from pathlib import Path
import pytest
from test_gaia import gaia_config, gaia_tasks
from evoagentx.compactflow.gaia import GaiaToolSession, GaiaUnavailable, GaiaBudgetExceeded, TOOL_SCHEMAS
from evoagentx.compactflow.gaia_contracts import CONTRACTS, contract_for, VERSION
from evoagentx.compactflow.gaia_tools import GaiaBackends
from evoagentx.compactflow.paper_workflow import compile_spec, validate_spec
from evoagentx.compactflow.runtime import CompactFlowRuntime
from evoagentx.compactflow.schema import Partial, Complete


def session(tmp_path,backend=None,config=None,output=None):
    c=config or gaia_config(tmp_path)
    task=gaia_tasks(tmp_path)[0]
    return GaiaToolSession(c,task.public_input(),output=output or tmp_path/'out',workspace=tmp_path/'work',seed=42,backend=backend)


def node(tool='search', fields=('first','second')):
    return {'id':'tools','tool':'gaia_'+tool,'instruction':'explicit observations',
            'inputs':{'arguments':'builder.arguments'},'outputs':list(fields)}

def args(**values):
    return {'arguments':json.dumps({'units':[{'field':k,'arguments':v} for k,v in values.items()]})}

class Backend:
    def __init__(self):self.calls=[]
    async def estimate(self,*a):return 0
    async def call(self,tool,arguments):
        self.calls.append((tool,arguments));await asyncio.sleep(.002)
        return {'value':arguments}


def test_registry_covers_every_tool_and_reports_provenance():
    assert set(CONTRACTS)==set(TOOL_SCHEMAS)|{'agent'}
    assert all(c.early_safe for c in CONTRACTS.values())
    assert all(c.effect_class!='pure' for c in CONTRACTS.values())
    assert all(c.metadata()['annotation_accuracy'] is None for c in CONTRACTS.values())
    with pytest.raises(ValueError):contract_for('gaia_unregistered')


def test_real_guarded_runtime_starts_tool_early_and_consumes_first_unit(tmp_path):
    async def run():
        tool_started=asyncio.Event();first_consumed=asyncio.Event()
        class B(Backend):
            async def call(self,tool,a):
                self.calls.append((tool,a));tool_started.set()
                if a['query']=='two':await asyncio.wait_for(first_consumed.wait(),2)
                return {'query':a['query']}
        backend=B();s=session(tmp_path,backend)
        class Client:
            async def stream(self,system,prompt,**kwargs):
                req=json.loads(prompt)
                if req['instruction']=='builder':
                    contract=req['gaia_output_contracts'][0]
                    assert contract['required_observation_fields']==['first','second']
                    assert contract['tool']=='gaia_search'
                    yield json.dumps({'field':'arguments','value':args(first={'query':'one'},second={'query':'two'})['arguments']})+'\n'
                    await asyncio.wait_for(tool_started.wait(),2)
                    yield json.dumps({'field':'tail','value':'done'})+'\n'
                else:
                    if req['instruction']=='reader':
                        assert json.loads(req['inputs']['value'])['query']=='one'
                        first_consumed.set()
                    yield json.dumps({'field':'answer','value':'ok'})+'\n'
        spec={'nodes':[
            {'id':'builder','tool':'llm','instruction':'builder','inputs':{'question':'$input.question'},'outputs':['arguments','tail']},
            node(),
            {'id':'reader','tool':'llm','instruction':'reader','inputs':{'value':'tools.first'},'outputs':['answer']},
            {'id':'sink','tool':'llm','instruction':'sink','inputs':{'answer':'reader.answer','second':'tools.second','tail':'builder.tail'},'outputs':['answer']}],
            'sinks':['sink'],'applied_policies':[],'unapplied_policies':[]}
        public=s.task;validate_spec(spec,public,set())
        from evoagentx.compactflow.replay import ReplayBundle,graph_to_dict,graph_from_replay
        from evoagentx.compactflow.execution_baselines import replay_llmcompiler,replay_llmorch
        bundle=ReplayBundle('gaia-units')
        graph=compile_spec(spec,public,Client(),seed=42,tool_session=s,recorder=bundle)
        assert ('builder','tools') in {(g.producer,g.consumer) for g in graph.guards}
        result=await CompactFlowRuntime(graph,mode='guarded',sink_ids=['sink'],call_timeout=3).execute(public)
        assert not result.errors,result.errors
        assert result.outputs['sink']['answer']=='ok' and len(backend.calls)==2
        commits=[json.loads(x) for x in (tmp_path/'out/contract_records.jsonl').read_text().splitlines()]
        assert [c['field'] for c in commits]==['first','second']
        assert all(c['observation_key'] and c['contract_version']==VERSION for c in commits)
        for mode in ['llmorch','llmcompiler','guarded']:
            replay_graph=graph_from_replay(graph_to_dict(graph,['sink']),bundle,time_scale=0)
            if mode=='guarded':
                replay=await CompactFlowRuntime(replay_graph,mode=mode,sink_ids=['sink']).execute(public)
            else:
                runner=replay_llmorch if mode=='llmorch' else replay_llmcompiler
                replay=await runner(replay_graph,public,sinks=('sink',),call_timeout=3,workflow_timeout=10,
                                    **({'kinds':{c.id:'inout' for c in replay_graph.calls},'processors':4} if mode=='llmorch' else {}))
            assert not replay.errors and replay.outputs['sink']==result.outputs['sink']
        assert len(backend.calls)==2
    asyncio.run(run())


def test_asset_effect_barriers_override_early_data_guards(tmp_path):
    s=session(tmp_path,Backend())
    spec={'nodes':[{'id':'a','tool':'gaia_read_file','instruction':'read','inputs':{'arguments':'$input.question'},'outputs':['arguments']},
                   {'id':'b','tool':'gaia_vision','instruction':'see','inputs':{'arguments':'a.arguments'},'outputs':['answer']}],
          'sinks':['b'],'applied_policies':[],'unapplied_policies':[]}
    g=compile_spec(spec,s.task,None,seed=42,tool_session=s)
    assert any(d.producer=='a' and d.consumer=='b' and not d.allow_partial for d in g.effect_dependencies)
    assert not any(d.producer=='a' and d.consumer=='b' for d in g.guards)
    assert all(c.metadata['effect_class']!='pure' for c in g.calls)


@pytest.mark.parametrize('units',[
    [{'field':'first','arguments':{}},{'field':'first','arguments':{}}],
    [{'field':'first','arguments':{}}],
    [{'field':'first','arguments':{'units':[]}},{'field':'second','arguments':{}}],
])
def test_bad_units_fail_before_any_tool_call(tmp_path,units):
    backend=Backend();s=session(tmp_path,backend)
    async def run():
        with pytest.raises(ValueError):
            async for _ in s.workflow_stream(None,node(),{'arguments':json.dumps({'units':units})}):pass
        assert not backend.calls and s.calls==0
    asyncio.run(run())


def test_partial_close_resume_and_exact_replay_do_not_repeat_tools(tmp_path):
    backend=Backend();s=session(tmp_path,backend);request=args(first={'query':'one'},second={'query':'two'})
    async def run():
        async with aclosing(s.workflow_stream(None,node(),request)) as stream:
            assert isinstance(await anext(stream),Partial)
        assert len(backend.calls)==1
        restored=session(tmp_path,backend,config=s.config)
        events=[e async for e in restored.workflow_stream(None,node(),request)]
        assert len(backend.calls)==2 and isinstance(events[-1],Complete)
        cfg=copy.deepcopy(s.config);cfg['tools']['gaia']['mode']='replay'
        class NoNetwork(Backend):
            async def call(self,*a):pytest.fail('replay invoked a live tool')
        replay=session(tmp_path,NoNetwork(),config=cfg)
        again=[e async for e in replay.workflow_stream(None,node(),request)]
        assert again[-1].data==events[-1].data
        with pytest.raises(GaiaUnavailable,match='missing'):
            async for _ in replay.workflow_stream(None,node(),args(first={'query':'different'},second={'query':'two'})):pass
        assert replay.incomplete
    asyncio.run(run())


def test_budget_exhaustion_preserves_first_commit_without_unknown_cost(tmp_path):
    backend=Backend();s=session(tmp_path,backend);s.cfg['max_tool_calls']=2
    async def run():
        await s.invoke('search',{'query':'prior'})
        events=[]
        with pytest.raises(GaiaBudgetExceeded):
            async for e in s.workflow_stream(None,node(),args(first={'query':'one'},second={'query':'two'})):events.append(e)
        assert len(events)==1 and isinstance(events[0],Partial)
        assert len(backend.calls)==2 and not s.incomplete
    asyncio.run(run())


def test_agent_final_is_complete_only_but_can_consume_stable_inputs(tmp_path):
    s=session(tmp_path,Backend())
    class Client:
        async def json(self,*a,**k):return {'final':{'answer':'done'}}
    async def run():
        events=[e async for e in s.workflow_stream(Client(),node('agent',('answer',)),{'question':'q'})]
        assert len(events)==1 and isinstance(events[0],Complete)
        graph=compile_spec({'nodes':[{'id':'agent','tool':'gaia_agent','instruction':'solve','inputs':{'question':'$input.question'},'outputs':['answer']}], 'sinks':['agent']},s.task,Client(),seed=42,tool_session=s)
        c=graph.calls[0]
        assert c.early_safe and c.stream_contract.stable_fields==()
        assert c.metadata['output_mode']=='complete_final_only'
    asyncio.run(run())


def test_native_adapters_use_the_same_effect_registry():
    from evoagentx.compactflow.baseline_workflows import lower_aflow,export_aflow
    from test_baselines import GRAPH
    artifact=export_aflow(GRAPH.replace('operator.Custom','operator.GAIAToolAgent'),"SOLVE='solve'")
    graph=lower_aflow(artifact,{'problem':'q','entry_point':None},None,4)
    assert graph.calls[0].metadata['contract_version']==VERSION
    assert graph.calls[0].metadata['footprint_source']=='static_export_full_result_bindings'
    assert graph.calls[0].effects


def test_document_units_and_asset_index_concurrency(tmp_path):
    from concurrent.futures import ThreadPoolExecutor
    from pypdf import PdfWriter
    from docx import Document
    from pptx import Presentation
    import openpyxl
    s=session(tmp_path)
    backend=s.backend
    pdf=tmp_path/'two.pdf';w=PdfWriter();w.add_blank_page(100,100);w.add_blank_page(100,100);w.write(pdf)
    doc=Document();doc.add_paragraph('first');doc.add_paragraph('second');doc.save(tmp_path/'two.docx')
    deck=Presentation()
    for text in ['first','second']:
        slide=deck.slides.add_slide(deck.slide_layouts[5]);slide.shapes.title.text=text
    deck.save(tmp_path/'two.pptx')
    book=openpyxl.Workbook();book.active.append(['first']);book.active.append(['second']);book.save(tmp_path/'two.xlsx')
    for name,unit,expected in [('two.pdf',{'kind':'pdf_page','index':2},''),('two.docx',{'kind':'paragraph','index':2},'second'),('two.pptx',{'kind':'slide','index':2},'second'),('two.xlsx',{'kind':'rows','start':2,'count':1},'second')]:
        asset=backend.save_asset((tmp_path/name).read_bytes(),name)
        value=backend.tool_read_file(asset['asset_id'],unit=unit)
        assert expected in value['text'] and value['unit']==unit
    other=GaiaBackends(s.config,s.task,backend.workspace)
    with ThreadPoolExecutor(2) as pool:
        a=pool.submit(backend.save_asset,b'new A','a.txt');b=pool.submit(other.save_asset,b'new B','b.txt')
        a,b=a.result(),b.result()
    index=json.loads(backend.index_path.read_text())
    assert a['asset_id'] in index and b['asset_id'] in index
    with pytest.raises(ValueError,match='not supplied'):backend.resolve(b['asset_id'])

@pytest.mark.parametrize('tool,usage', [
    ('vision',{'total_tokens':7,'usage_complete':True,'prompt_tokens':5,'completion_tokens':2}),
    ('transcribe',{'total_tokens':0,'usage_complete':True,'audio_seconds':1.25}),
])
def test_auxiliary_units_account_logical_and_physical_costs(tmp_path,tool,usage):
    class Aux(Backend):
        async def call(self,*a):
            self.calls.append(a)
            return {'text':'observed','_usage':copy.deepcopy(usage)}
    backend=Aux();s=session(tmp_path,backend)
    request=args(first={'asset_id':'one'},second={'asset_id':'two'})
    async def run():
        events=[e async for e in s.workflow_stream(None,node(tool),request)]
        assert isinstance(events[-1],Complete)
        restored=session(tmp_path,backend,config=s.config)
        again=[e async for e in restored.workflow_stream(None,node(tool),request)]
        assert again[-1].data==events[-1].data and len(backend.calls)==2
        metric='vision_tokens' if tool=='vision' else 'audio_seconds'
        physical='physical_'+metric
        expected=14 if tool=='vision' else 2.5
        assert s.accounting()[metric]==expected and s.accounting()[physical]==expected
        assert restored.accounting()[metric]==expected and restored.accounting()[physical]==0
    asyncio.run(run())

@pytest.mark.parametrize('tool', ['vision','transcribe'])
def test_missing_auxiliary_usage_cannot_commit_stable_field(tmp_path,tool):
    s=session(tmp_path,Backend())
    async def run():
        with pytest.raises(GaiaUnavailable,match='measured'):
            async for _ in s.workflow_stream(None,node(tool,('answer',)),{'asset_id':'one'}):
                pytest.fail('unknown auxiliary cost produced a stable observation')
        assert s.incomplete and not (tmp_path/'out/contract_records.jsonl').exists()
    asyncio.run(run())


def test_cancellation_after_first_unit_never_retries_unresolved_second(tmp_path):
    async def run():
        started=asyncio.Event()
        class Slow(Backend):
            async def call(self,tool,a):
                self.calls.append((tool,a))
                if a['query']=='two':
                    started.set();await asyncio.Event().wait()
                return a
        backend=Slow();s=session(tmp_path,backend)
        request=args(first={'query':'one'},second={'query':'two'})
        stream=s.workflow_stream(None,node(),request)
        assert isinstance(await anext(stream),Partial)
        task=asyncio.create_task(anext(stream));await started.wait();task.cancel()
        with pytest.raises(asyncio.CancelledError):await task
        assert s.incomplete and s.records[-1]['status']=='incomplete'
        restored=session(tmp_path,backend,config=s.config)
        with pytest.raises(GaiaUnavailable,match='unresolved'):
            async for _ in restored.workflow_stream(None,node(),request):pass
        assert len(backend.calls)==2
    asyncio.run(run())


def test_target_gate_blocks_primitive_and_agent_before_calls(tmp_path):
    s=session(tmp_path,Backend());s.task['split']='target'
    class NoModel:
        async def json(self,*a,**k):pytest.fail('target model called before freeze')
    async def run():
        for n in [node(),node('agent',('answer',))]:
            with pytest.raises(GaiaUnavailable,match='frozen'):
                async for _ in s.workflow_stream(NoModel(),n,{}):pass
        assert s.calls==0 and not s.backend.calls
    asyncio.run(run())


def test_primitive_tools_remain_gaia_only(tmp_path):
    s=session(tmp_path,Backend())
    spec={'nodes':[{'id':'tool','tool':'gaia_search','instruction':'search','inputs':{'arguments':'$input.question'},'outputs':['answer']}],
          'sinks':['tool'],'applied_policies':[],'unapplied_policies':[]}
    validate_spec(spec,s.task,set())
    for benchmark in ['MBPP','MATH','HotpotQA']:
        with pytest.raises(ValueError,match='unregistered'):validate_spec(spec,{**s.task,'benchmark':benchmark},set())


def test_zip_members_audio_windows_and_video_frames(tmp_path):
    import io,zipfile,wave,subprocess
    s=session(tmp_path);b=s.backend
    data=io.BytesIO()
    with zipfile.ZipFile(data,'w') as z:
        z.writestr('a.txt','aaa');z.writestr('b.txt','bbb')
    asset=b.save_asset(data.getvalue(),'two.zip')
    result=b.tool_unpack_zip(asset['asset_id'],members=['b.txt'])
    assert len(result['assets'])==1 and result['assets'][0]['name']=='b.txt'
    with pytest.raises(ValueError):b.tool_unpack_zip(asset['asset_id'],members=['../a.txt'])
    audio=io.BytesIO()
    with wave.open(audio,'wb') as w:
        w.setnchannels(1);w.setsampwidth(2);w.setframerate(16000);w.writeframes(b'\0\0'*32000)
    asset=b.save_asset(audio.getvalue(),'silence.wav')
    def aux(endpoint,payload):
        with wave.open(payload['path'],'rb') as w:duration=w.getnframes()/w.getframerate()
        return {'text':'','_usage':{'usage_complete':True,'audio_seconds':duration}}
    b.aux=aux
    result=b.tool_transcribe(asset['asset_id'],start_seconds=.5,seconds=.75)
    assert result['_usage']['audio_seconds']==.75
    video=tmp_path/'tiny.mp4'
    subprocess.run(['ffmpeg','-nostdin','-v','error','-f','lavfi','-i','color=c=red:s=32x32:d=1','-c:v','mpeg4',str(video)],check=True)
    asset=b.save_asset(video.read_bytes(),'tiny.mp4')
    result=b.tool_video_frames(asset['asset_id'],frame_seconds=.2)
    assert len(result['assets'])==1 and result['requested_frame_seconds']==.2


def test_planner_receives_primitive_binding_example_and_rejects_misbound_args(tmp_path):
    from evoagentx.compactflow.paper_workflow import plan_workflow
    cfg=gaia_config(tmp_path);task=gaia_tasks(tmp_path)[0]
    class Planner:
        async def json(self,system,prompt,**kwargs):
            request=json.loads(prompt)
            example=request['gaia_primitive_example']
            assert example['nodes'][1]['inputs']=={'arguments':'build_args.arguments'}
            assert 'gaia_read_file' in request['workflow_schema']['properties']['nodes']['items']['properties']['tool']['enum']
            return example
    spec,_=asyncio.run(plan_workflow(Planner(),task,[],cfg['construction']['planner'],seed=42))
    bad=copy.deepcopy(spec);bad['nodes'][1]['inputs']={'asset_id':'$input.assets','unit':'rows.sheet'}
    with pytest.raises(ValueError,match='INSIDE that JSON'):
        validate_spec(bad,task.public_input(),set())
