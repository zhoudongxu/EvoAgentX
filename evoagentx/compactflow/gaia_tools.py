"""Concrete GAIA tools; execution never downloads models or browser drivers."""

from __future__ import annotations
import asyncio, ipaddress, json, re, socket, stat, subprocess, tempfile, urllib.parse, urllib.request, zipfile
import threading
import fcntl
import os
import math
from pathlib import Path
from .gaia import ASSET_EXTENSIONS, GaiaUnavailable, file_sha
from .evolution import atomic_json
from .replay import canonical

_TOOL_CAPACITY = threading.BoundedSemaphore(4)


def public_url(url):
    p = urllib.parse.urlsplit(url)
    if p.scheme not in {"http", "https"} or not p.hostname or p.username or p.password:
        raise ValueError("only public HTTP(S) URLs are accepted")
    if any(
        s in urllib.parse.unquote(url).lower()
        for s in (
            "gaia-benchmark",
            "/gaia/2023/",
            "metadata.parquet",
            "annotator_metadata",
        )
    ):
        raise ValueError(
            "benchmark answers/annotations are prohibited retrieval sources"
        )
    for item in socket.getaddrinfo(
        p.hostname,
        p.port or (443 if p.scheme == "https" else 80),
        type=socket.SOCK_STREAM,
    ):
        if not ipaddress.ip_address(item[4][0]).is_global:
            raise ValueError("non-public data-source address")
    return url


class PublicRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        public_url(newurl)
        return super().redirect_request(req, fp, code, msg, headers, newurl)


class GaiaBackends:
    def __init__(self, config, task, workspace):
        self.config, self.cfg, self.task = config, config["tools"]["gaia"], task
        self.workspace = Path(workspace)
        self.files = self.workspace / "files"
        self.files.mkdir(parents=True, exist_ok=True)
        self.index_path = self.workspace / "index.json"
        self.assets = (
            json.loads(self.index_path.read_text()) if self.index_path.exists() else {}
        )
        # Physical blobs may be shared by methods; visible assets are session-local.
        self.authorized = set()
        root = Path(self.cfg["asset_root"]).resolve()
        for a in task.get("assets", []):
            source = (root / a["name"]).resolve()
            if (
                not source.is_relative_to(root)
                or not source.is_file()
                or file_sha(source) != a["sha256"]
            ):
                raise GaiaUnavailable(
                    "attachment missing/checksum changed: " + a["name"]
                )
            self.save_asset(source.read_bytes(), a["name"], expected=a["asset_id"])

    def save_asset(self, content, name, expected=None):
        import hashlib

        safe = Path(name).name
        ext = Path(safe).suffix.lower()
        if ext not in ASSET_EXTENSIONS or safe.lower().startswith("metadata."):
            raise ValueError("unsupported/prohibited asset type: " + ext)
        if len(content) > self.cfg.get("max_asset_bytes", 100 * 1024 * 1024):
            raise ValueError("asset exceeds byte limit")
        sha = hashlib.sha256(content).hexdigest()
        key, filename = "sha256:" + sha, sha + ext
        if expected and key != expected:
            raise GaiaUnavailable("asset identity mismatch")
        path = self.files / filename
        value = dict(asset_id=key,name=safe,file=filename,sha256=sha,bytes=len(content),type=ext)
        # Multiple ready nodes/methods may commit assets concurrently. Merge the
        # content-addressed index under a process lock; never overwrite a sibling.
        with (self.workspace / "index.lock").open("a") as guard:
            fcntl.flock(guard, fcntl.LOCK_EX)
            if path.exists() and file_sha(path) != sha:
                raise GaiaUnavailable("stored asset was modified")
            if not path.exists():
                fd, temp = tempfile.mkstemp(dir=self.files)
                try:
                    with os.fdopen(fd,"wb") as f:
                        f.write(content); f.flush(); os.fsync(f.fileno())
                    os.replace(temp,path)
                finally:
                    Path(temp).unlink(missing_ok=True)
            merged = json.loads(self.index_path.read_text()) if self.index_path.exists() else {}
            # The same bytes may be discovered under another name; retain the
            # first immutable descriptor rather than rewriting prior observations.
            value = merged.setdefault(key,value)
            atomic_json(self.index_path,merged)
            self.assets.update(merged)
            self.authorized.add(key)
        return {k: v for k, v in value.items() if k != "file"}

    def resolve(self, asset_id):
        if asset_id not in self.authorized:
            raise ValueError("asset was not supplied to or observed by this session")
        if asset_id not in self.assets:
            self.assets = (
                json.loads(self.index_path.read_text())
                if self.index_path.exists()
                else {}
            )
        if asset_id not in self.assets:
            raise ValueError("unknown asset_id; only this task assets are accessible")
        entry = self.assets[asset_id]
        path = (self.files / entry["file"]).resolve()
        if (
            not path.is_relative_to(self.files.resolve())
            or not path.is_file()
            or file_sha(path) != entry["sha256"]
        ):
            raise GaiaUnavailable("task asset checksum/path mismatch")
        return path

    def verify_result_assets(self, value):
        if isinstance(value, dict):
            if "asset_id" in value:
                # A verified, exactly matching cached observation grants this asset.
                self.authorized.add(value["asset_id"])
                self.resolve(value["asset_id"])
            for v in value.values():
                self.verify_result_assets(v)
        elif isinstance(value, list):
            for v in value:
                self.verify_result_assets(v)

    def fetch(self, url):
        public_url(url)
        req = urllib.request.Request(
            url, headers={"User-Agent": "CompactFlow-GAIA/1.0 (research; read-only)"}
        )
        with urllib.request.build_opener(PublicRedirect()).open(
            req, timeout=45
        ) as response:
            public_url(response.url)
            limit = self.cfg.get("max_asset_bytes", 100 * 1024 * 1024)
            body = response.read(limit + 1)
            if len(body) > limit:
                raise ValueError("download exceeds byte limit")
            return body, response.headers.get_content_type(), response.url

    def aux(self, endpoint, payload):
        req = urllib.request.Request(
            self.cfg["auxiliary_url"].rstrip("/") + "/" + endpoint,
            data=json.dumps(payload).encode(),
            headers={"Content-Type": "application/json"},
        )
        try:
            with urllib.request.urlopen(
                req, timeout=self.cfg.get("tool_timeout_seconds", 180)
            ) as r:
                value = json.load(r)
        except Exception as exc:
            raise GaiaUnavailable("auxiliary request failed: " + str(exc)) from exc
        if value.get("error"):
            raise GaiaUnavailable(value["error"])
        return value

    def vision_payload(self, args):
        path = self.resolve(args["asset_id"])
        if path.suffix.lower() not in {".png", ".jpg", ".jpeg"}:
            raise ValueError("vision requires an image; render documents/videos first")
        return dict(
            path=str(path),
            sha256=file_sha(path),
            question=str(args["question"]),
            max_tokens=self.cfg.get("vision_max_tokens", 1024),
        )

    async def estimate(self, tool, args):
        if tool != "vision":
            return 0
        payload = self.vision_payload(args)
        v = await asyncio.to_thread(self.aux, "estimate", payload)
        return int(v["prompt_tokens"]) + payload["max_tokens"]

    async def call(self, tool, args):
        fn = getattr(self, "tool_" + tool, None)
        if fn is None:
            raise ValueError("unregistered backend")

        def bounded():
            # Held by the worker until it actually finishes, even after cancellation.
            with _TOOL_CAPACITY:
                return fn(**args)

        return await asyncio.to_thread(bounded)

    def tool_read_file(self, asset_id, offset=0, limit=12000, page=None, unit=None):
        path = self.resolve(asset_id)
        offset, limit = int(offset), int(limit)
        if offset < 0 or not 1 <= limit <= 24000:
            raise ValueError("invalid read offset/limit")
        ext, assets, info = path.suffix.lower(), [], {}
        if unit is not None:
            if page is not None:
                raise ValueError("page rendering and a text unit are separate requests")
            text, info = self.read_unit(path, unit)
            return dict(asset_id=asset_id,unit=unit,text=text[offset:offset+limit],offset=offset,
                        total_characters=len(text),next_offset=offset+limit if offset+limit<len(text) else None,
                        assets=[],**info)
        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(path)
            texts = [p.extract_text() or "" for p in reader.pages]
            info["pages"] = len(texts)
            text = "\n".join(f"PAGE {i + 1}\n{t}" for i, t in enumerate(texts))
            if page is not None or not any(t.strip() for t in texts):
                page = int(page or 1)
                if not 1 <= page <= len(texts):
                    raise ValueError("PDF page outside file")
                with tempfile.TemporaryDirectory(dir=self.workspace) as d:
                    dest = Path(d) / "page"
                    subprocess.run(
                        [
                            "pdftoppm",
                            "-f",
                            str(page),
                            "-l",
                            str(page),
                            "-singlefile",
                            "-scale-to",
                            "1600",
                            "-png",
                            str(path),
                            str(dest),
                        ],
                        check=True,
                        capture_output=True,
                        timeout=60,
                    )
                    assets.append(
                        self.save_asset(
                            dest.with_suffix(".png").read_bytes(), f"page-{page}.png"
                        )
                    )
        elif ext == ".xlsx":
            import openpyxl

            book = openpyxl.load_workbook(path, read_only=True, data_only=True)
            parts = []
            for sheet in book:
                parts.append("SHEET " + sheet.title)
                for i, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    parts.append(
                        canonical(
                            {
                                "row": i,
                                "values": [
                                    str(v) if v is not None else None for v in row
                                ],
                            }
                        )
                    )
            book.close()
            text = "\n".join(parts)
        elif ext == ".xls":
            import xlrd

            book = xlrd.open_workbook(str(path), on_demand=True)
            parts = []
            for sheet in book.sheets():
                parts.append("SHEET " + sheet.name)
                for index in range(sheet.nrows):
                    parts.append(
                        canonical({"row": index + 1, "values": sheet.row_values(index)})
                    )
            book.release_resources()
            text = "\n".join(parts)
        elif ext == ".docx":
            from docx import Document

            doc = Document(path)
            text = "\n".join(
                [p.text for p in doc.paragraphs]
                + [
                    canonical([[c.text for c in r.cells] for r in t.rows])
                    for t in doc.tables
                ]
            )
            assets = self.embedded_images(path)
        elif ext == ".pptx":
            from pptx import Presentation

            text = "\n".join(
                f"SLIDE {i + 1}\n"
                + "\n".join(s.text for s in slide.shapes if hasattr(s, "text"))
                for i, slide in enumerate(Presentation(path).slides)
            )
            assets = self.embedded_images(path)
        elif ext in {
            ".txt",
            ".csv",
            ".json",
            ".jsonld",
            ".xml",
            ".pdb",
            ".py",
            ".html",
        }:
            text = path.read_text(encoding="utf-8", errors="replace")
        else:
            raise ValueError(
                "use unpack_zip, vision, transcribe, or video_frames for this type"
            )
        return dict(
            asset_id=asset_id,
            text=text[offset : offset + limit],
            offset=offset,
            total_characters=len(text),
            next_offset=offset + limit if offset + limit < len(text) else None,
            assets=assets,
            **info,
        )

    def read_unit(self, path, unit):
        """Read only the requested immutable content unit; do not pre-extract later units."""
        if not isinstance(unit,dict) or "kind" not in unit:
            raise ValueError("unit requires an explicit kind")
        kind = unit["kind"]
        allowed = {"kind","index"} if kind in {"pdf_page","slide","paragraph"} else ({"kind","sheet","start","count"} if kind=="rows" else {"kind"})
        if set(unit)-allowed:raise ValueError("unknown unit selector fields")
        def integer(name,default):
            v=unit.get(name,default)
            if type(v) is not int or v<1:raise ValueError("unit indices/counts must be positive integers")
            return v
        ext=path.suffix.lower()
        if kind=="pdf_page" and ext==".pdf":
            from pypdf import PdfReader
            reader=PdfReader(path);i=integer("index",1)
            if i>len(reader.pages):raise ValueError("PDF page outside file")
            return reader.pages[i-1].extract_text() or "", {"page":i,"pages":len(reader.pages)}
        if kind=="slide" and ext==".pptx":
            from pptx import Presentation
            slides=Presentation(path).slides;i=integer("index",1)
            if i>len(slides):raise ValueError("slide outside presentation")
            return "\n".join(x.text for x in slides[i-1].shapes if hasattr(x,"text")), {"slide":i}
        if kind=="paragraph" and ext==".docx":
            from docx import Document
            paragraphs=Document(path).paragraphs;i=integer("index",1)
            if i>len(paragraphs):raise ValueError("paragraph outside document")
            return paragraphs[i-1].text, {"paragraph":i}
        if kind=="rows" and ext in {".xlsx",".xls",".csv"}:
            start=integer("start",1);count=integer("count",100)
            if count>1000:raise ValueError("row unit exceeds 1000 rows")
            sheet=unit.get("sheet")
            if ext==".xlsx":
                import openpyxl
                book=openpyxl.load_workbook(path,read_only=True,data_only=True)
                try:
                    ws=book[sheet] if sheet is not None else book.worksheets[0]
                    rows=list(ws.iter_rows(min_row=start,max_row=min(start+count-1,ws.max_row),values_only=True)) if start<=ws.max_row else []
                finally:book.close()
            elif ext==".xls":
                import xlrd
                book=xlrd.open_workbook(str(path),on_demand=True)
                try:
                    ws=book.sheet_by_name(sheet) if sheet is not None else book.sheet_by_index(0)
                    rows=[ws.row_values(i) for i in range(start-1,min(start+count-1,ws.nrows))]
                finally:book.release_resources()
            else:
                import csv,itertools
                if sheet is not None:raise ValueError("CSV has no named sheet")
                with path.open(encoding="utf-8",errors="replace",newline="") as f:
                    rows=list(itertools.islice(csv.reader(f),start-1,start+count-1))
            if not rows:raise ValueError("row unit outside table")
            return "\n".join(canonical({"row":start+i,"values":[str(v) if v is not None else None for v in row]}) for i,row in enumerate(rows)), {"start_row":start,"rows":len(rows)}
        if kind=="text" and ext in {".txt",".csv",".json",".jsonld",".xml",".pdb",".py",".html"}:
            return path.read_text(encoding="utf-8",errors="replace"), {}
        raise ValueError("unit kind does not match the authorized document format")

    def embedded_images(self, path):
        assets = []
        with zipfile.ZipFile(path) as archive:
            for item in archive.infolist():
                if "/media/" in item.filename and Path(
                    item.filename
                ).suffix.lower() in {".png", ".jpg", ".jpeg"}:
                    if item.file_size > self.cfg.get(
                        "max_asset_bytes", 100 * 1024 * 1024
                    ):
                        raise ValueError("embedded image too large")
                    assets.append(
                        self.save_asset(archive.read(item), Path(item.filename).name)
                    )
        return assets

    def tool_unpack_zip(self, asset_id, members=None):
        path = self.resolve(asset_id)
        if path.suffix != ".zip":
            raise ValueError("unpack_zip requires ZIP")
        with zipfile.ZipFile(path) as archive:
            entries = [x for x in archive.infolist() if not x.is_dir()]
            if len(entries) > 100 or sum(x.file_size for x in entries) > self.cfg.get(
                "max_asset_bytes", 100 * 1024 * 1024
            ):
                raise ValueError("archive expansion exceeds limits")
            for x in entries:
                normalized = x.filename.replace("\\", "/")
                if (
                    Path(normalized).is_absolute()
                    or ".." in Path(normalized).parts
                    or ":" in normalized
                    or stat.S_ISLNK(x.external_attr >> 16)
                ):
                    raise ValueError("unsafe ZIP member")
            if members is not None:
                if not isinstance(members,list) or not members or any(not isinstance(n,str) for n in members) or len(members)!=len(set(members)):
                    raise ValueError("members must be distinct exact ZIP names")
                if set(members)-{x.filename for x in entries}:raise ValueError("unknown ZIP member")
                entries=[x for x in entries if x.filename in members]
            return {
                "assets": [
                    self.save_asset(archive.read(x), Path(x.filename).name)
                    for x in entries
                ]
            }

    def tool_search(self, query):
        if re.search(r"gaia.{0,30}(benchmark|dataset|answer|solution)", query, re.I):
            raise ValueError("answer-key query prohibited")
        backend = self.cfg.get("search_backend", "duckduckgo")
        if backend == "bing_html":
            from evoagentx.tools.search_base import SearchBase
            from bs4 import BeautifulSoup
            import base64

            url = "https://www.bing.com/search?" + urllib.parse.urlencode(
                {"q": str(query), "count": 5}
            )
            body, content_type, final = self.fetch(url)
            soup = BeautifulSoup(body, "html.parser")
            formatter = SearchBase(num_search_pages=5, max_content_words=350)
            results = []
            for item in soup.select("li.b_algo"):
                link = item.select_one("h2 a[href]")
                if link is None:
                    continue
                address = link["href"]
                parsed = urllib.parse.urlsplit(address)
                if (
                    parsed.hostname
                    and parsed.hostname.endswith("bing.com")
                    and parsed.path.startswith("/ck/")
                ):
                    encoded = urllib.parse.parse_qs(parsed.query).get("u", [""])[0]
                    if encoded.startswith("a1"):
                        try:
                            address = base64.urlsafe_b64decode(
                                encoded[2:] + "=" * (-len(encoded[2:]) % 4)
                            ).decode()
                        except (ValueError, UnicodeDecodeError):
                            continue
                try:
                    public_url(address)
                except ValueError:
                    continue
                results.append(
                    {
                        "title": link.get_text(" ", strip=True),
                        "url": address,
                        "content": formatter._truncate_content(
                            item.get_text(" ", strip=True), 350
                        ),
                    }
                )
                if len(results) == 5:
                    break
            if not results:
                raise ValueError(
                    "fixed Bing search returned no usable results; revise the query or use browse"
                )
            return {"backend": backend, "source_url": final, "results": results}
        from ddgs.engines import ENGINES

        if backend not in ENGINES.get("text", {}):
            raise GaiaUnavailable(
                "fixed search backend unavailable; automatic provider fallback is prohibited"
            )
        from evoagentx.tools.search_ddgs import SearchDDGS

        parent = self

        class ControlledSearch(SearchDDGS):
            def _scrape_page(self, url):
                from bs4 import BeautifulSoup

                body, ctype, final = parent.fetch(url)
                if "html" not in ctype:
                    return "", ""
                soup = BeautifulSoup(body, "html.parser")
                return (soup.title.get_text() if soup.title else final), soup.get_text(
                    " ", strip=True
                )[:12000]

        result = ControlledSearch(
            backend=self.cfg.get("search_backend", "duckduckgo"),
            num_search_pages=5,
            max_content_words=1500,
        ).search(str(query))
        if result.get("error"):
            raise GaiaUnavailable("search unavailable: " + str(result["error"]))
        kept = []
        for item in result.get("results", []):
            try:
                public_url(item.get("url", item.get("href", "")))
            except ValueError:
                continue
            kept.append(item)
        return {
            "results": kept,
            "backend": self.cfg.get("search_backend", "duckduckgo"),
        }

    def tool_browse(self, url, offset=0, limit=12000):
        public_url(url)
        from evoagentx.tools.browser_tool import BrowserBase
        from selenium import webdriver
        from selenium.webdriver.firefox.options import Options
        from selenium.webdriver.firefox.service import Service
        from bs4 import BeautifulSoup

        opts = Options()
        opts.add_argument("-headless")
        if self.cfg.get("browser_binary"):
            opts.binary_location = self.cfg["browser_binary"]
        browser = BrowserBase(browser_type="firefox", headless=True, timeout=45)
        try:
            browser.driver = webdriver.Firefox(
                service=Service(self.cfg["driver_path"]), options=opts
            )
            browser.driver.set_page_load_timeout(45)
            result = browser.navigate_to_url(url)
            if result.get("status") == "error":
                raise GaiaUnavailable(str(result))
            final = public_url(browser.driver.current_url)
            soup = BeautifulSoup(browser.driver.page_source, "html.parser")
            for tag in soup(["script", "style"]):
                tag.decompose()
            text = soup.get_text(" ", strip=True)
            offset, limit = int(offset), min(24000, int(limit))
            if offset < 0 or limit <= 0:
                raise ValueError("invalid browser offset/limit")
            links = [
                {
                    "text": a.get_text(" ", strip=True),
                    "url": urllib.parse.urljoin(final, a["href"]),
                }
                for a in soup.select("a[href]")
            ][:100]
            shot = self.save_asset(browser.driver.get_screenshot_as_png(), "page.png")
            return dict(
                url=final,
                title=browser.driver.title,
                text=text[offset : offset + limit],
                next_offset=offset + limit if len(text) > offset + limit else None,
                links=links,
                assets=[shot],
            )
        finally:
            browser.close_browser()

    def tool_download(self, url, media=False):
        if media:
            public_url(url)
            import yt_dlp

            class PublicMedia(yt_dlp.YoutubeDL):
                def urlopen(self, req):
                    address = (
                        req
                        if isinstance(req, str)
                        else getattr(req, "url", getattr(req, "full_url", None))
                    )
                    public_url(address)
                    return super().urlopen(req)

            with tempfile.TemporaryDirectory(dir=self.workspace) as d:
                options = {
                    "quiet": True,
                    "noplaylist": True,
                    "socket_timeout": 30,
                    "max_filesize": self.cfg.get("max_asset_bytes", 100 * 1024 * 1024),
                    "format": "bestvideo[height<=720][ext=mp4]+bestaudio[ext=m4a]/best[height<=720][ext=mp4]/best[height<=720]",
                    "merge_output_format": "mp4",
                    "outtmpl": str(Path(d) / "media.%(ext)s"),
                    "max_downloads": 1,
                    "retries": 0,
                    "fragment_retries": 0,
                    "skip_unavailable_fragments": False,
                }
                def bounded_media(info, *, incomplete=False):
                    if info.get("is_live") or (info.get("duration") or 0) > 3600:
                        return "live media or media over one hour is outside the bounded tool"
                    return None
                options["match_filter"] = bounded_media
                try:
                    with PublicMedia(options) as downloader:
                        info = downloader.extract_info(url, download=True)
                except yt_dlp.utils.DownloadError as exc:
                    text = str(exc)
                    ordinary = ("requested format is not available", "video unavailable", "private video",
                                "video has been removed", "sign in to confirm", "http error 403", "http error 404", "maxdownloadsreached", "maximum number of downloads reached",
                                "members-only", "not available in your country")
                    if any(marker in text.lower() for marker in ordinary):
                        raise ValueError("media resource unavailable: " + text) from exc
                    raise
                if not info:
                    raise ValueError("media is outside the bounded download policy")
                paths = [
                    p
                    for p in Path(d).iterdir()
                    if p.suffix.lower() in {".mp4", ".webm", ".m4a", ".mp3"}
                ]
                if len(paths) != 1:
                    raise GaiaUnavailable("media export unavailable or incomplete")
                return {
                    "url": url,
                    "duration": info.get("duration"),
                    "assets": [self.save_asset(paths[0].read_bytes(), paths[0].name)],
                }
        body, ctype, final = self.fetch(url)
        name = Path(urllib.parse.urlsplit(final).path).name
        if Path(name).suffix.lower() not in ASSET_EXTENSIONS:
            ext = {
                "text/html": ".html",
                "image/jpeg": ".jpg",
                "image/png": ".png",
                "application/pdf": ".pdf",
                "audio/mpeg": ".mp3",
                "video/mp4": ".mp4",
            }.get(ctype)
            if not ext:
                raise ValueError("unsupported content type: " + ctype)
            name = "download" + ext
        return {"url": final, "assets": [self.save_asset(body, name)]}

    def tool_vision(self, asset_id, question):
        return self.aux(
            "vision", self.vision_payload(dict(asset_id=asset_id, question=question))
        )

    def tool_transcribe(self, asset_id, start_seconds=None, seconds=None):
        path = self.resolve(asset_id)
        if path.suffix.lower() not in {
            ".mp3",
            ".wav",
            ".m4a",
            ".ogg",
            ".flac",
            ".mp4",
            ".webm",
        }:
            raise ValueError("transcribe requires audio/video")
        if start_seconds is None and seconds is None:
            return self.aux("transcribe", dict(path=str(path), sha256=file_sha(path)))
        start=float(start_seconds or 0);duration=float(seconds or 0)
        if not math.isfinite(start) or not math.isfinite(duration) or start<0 or not 0<duration<=120:
            raise ValueError("audio window requires nonnegative start and 0 < seconds <= 120")
        with tempfile.TemporaryDirectory(dir=self.workspace) as d:
            clip=Path(d)/"segment.wav"
            subprocess.run(["ffmpeg","-nostdin","-v","error","-ss",str(start),"-i",str(path),"-t",str(duration),"-ac","1","-ar","16000",str(clip)],check=True,capture_output=True,timeout=90)
            result=self.aux("transcribe",dict(path=str(clip),sha256=file_sha(clip)))
            result["window"]={"start_seconds":start,"requested_seconds":duration}
            return result

    def tool_video_frames(self, asset_id, start_seconds=0, seconds=16, frame_seconds=None):
        path = self.resolve(asset_id)
        if path.suffix.lower() not in {".mp4", ".webm"}:
            raise ValueError("video asset required")
        if frame_seconds is not None:
            position=float(frame_seconds)
            if not math.isfinite(position) or position<0:raise ValueError("invalid frame timestamp")
            with tempfile.TemporaryDirectory(dir=self.workspace) as d:
                frame=Path(d)/"frame.jpg"
                subprocess.run(["ffmpeg","-nostdin","-v","error","-ss",str(position),"-i",str(path),"-vf","scale=1280:-2","-frames:v","1",str(frame)],check=True,capture_output=True,timeout=90)
                if not frame.exists():raise ValueError("frame outside video")
                return {"requested_frame_seconds":position,"assets":[self.save_asset(frame.read_bytes(),"frame.jpg")]}
        start_seconds, seconds = float(start_seconds), float(seconds)
        if start_seconds < 0 or not 0 < seconds <= 32:
            raise ValueError("invalid time window")
        with tempfile.TemporaryDirectory(dir=self.workspace) as d:
            subprocess.run(
                [
                    "ffmpeg",
                    "-nostdin",
                    "-v",
                    "error",
                    "-ss",
                    str(start_seconds),
                    "-i",
                    str(path),
                    "-t",
                    str(seconds),
                    "-vf",
                    "fps=1,scale=1280:-2",
                    "-frames:v",
                    "32",
                    str(Path(d) / "frame-%03d.jpg"),
                ],
                check=True,
                capture_output=True,
                timeout=90,
            )
            return {
                "start_seconds": start_seconds,
                "fps": 1,
                "assets": [
                    self.save_asset(p.read_bytes(), p.name)
                    for p in sorted(Path(d).glob("*.jpg"))
                ],
            }

    def tool_python(self, code):
        import uuid
        import shutil

        name = "compactflow-gaia-" + uuid.uuid4().hex
        image = self.config["tools"]["mbpp_sandbox"]["image"]
        if "@sha256:" not in image:
            raise GaiaUnavailable("sandbox image must be immutable")
        mount = tempfile.TemporaryDirectory(dir=self.workspace)
        Path(mount.name).chmod(0o755)
        visible = [self.assets[k] for k in sorted(self.authorized)]
        for entry in visible:
            dest = Path(mount.name) / entry["file"]
            shutil.copyfile(self.resolve(entry["asset_id"]), dest)
            dest.chmod(0o444)
        args = [
            "docker",
            "run",
            "--rm",
            "--name",
            name,
            "--network=none",
            "--read-only",
            "--cap-drop=ALL",
            "--security-opt=no-new-privileges",
            "--memory=256m",
            "--cpus=1",
            "--pids-limit=64",
            "--user=65534:65534",
            "--tmpfs=/tmp:rw,noexec,size=16m",
            "--mount",
            f"type=bind,src={mount.name},dst=/assets,readonly",
            "-i",
            image,
            "python",
            "-I",
            "-",
        ]
        listing = {v["name"]: "/assets/" + v["file"] for v in visible}
        try:
            result = subprocess.run(
                args,
                input="ASSETS = " + repr(listing) + "\n" + str(code),
                text=True,
                capture_output=True,
                timeout=20,
            )
            if result.returncode == 125:
                raise GaiaUnavailable("Docker unavailable: " + result.stderr[:1000])
            return dict(
                exit_code=result.returncode,
                stdout=result.stdout[:24000],
                stderr=result.stderr[:12000],
                truncated=len(result.stdout) > 24000 or len(result.stderr) > 12000,
                asset_paths=listing,
            )
        except subprocess.TimeoutExpired:
            return {"error": "Python exceeded 20 second limit"}
        finally:
            subprocess.run(
                ["docker", "rm", "-f", name], capture_output=True, timeout=15
            )
            mount.cleanup()
